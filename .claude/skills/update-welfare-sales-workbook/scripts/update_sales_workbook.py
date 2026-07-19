#!/usr/bin/env python3
"""Create or update a linked welfare-equipment sales workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


SHEETS = ["사업소_목록", "영업일지", "후속조치", "품목_관심도", "주간_현황", "확인필요"]
OFFICE_HEADERS = ["사업소ID", "사업소명", "지역", "담당자", "진행단계", "관심품목", "최근활동일", "다음행동", "다음행동일", "누적활동수", "등록일"]
LOG_HEADERS = ["활동ID", "활동일", "담당자", "사업소ID", "사업소명", "지역", "접촉방식", "상담내용", "관심품목", "진행단계", "다음행동", "예정일"]
FOLLOW_HEADERS = ["활동ID", "사업소명", "담당자", "다음행동", "예정일", "상태", "근거활동일"]
REVIEW_HEADERS = ["유형", "활동ID", "사업소명", "내용", "권장조치"]

NAVY = "17324D"
BLUE = "2563EB"
CYAN = "0EA5A8"
PALE = "EAF2F8"
YELLOW = "FFF4CC"
RED = "FDE2E2"
GREEN = "DCFCE7"
WHITE = "FFFFFF"
GRAY = "64748B"
THIN = Side(style="thin", color="D9E2EC")


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def norm_name(value):
    return re.sub(r"[^0-9a-z가-힣]", "", clean(value).lower())


def split_products(value):
    if isinstance(value, list):
        values = value
    else:
        values = re.split(r"[,;/|]", clean(value)) if value else []
    return sorted({clean(v) for v in values if clean(v)})


def stable_id(entry):
    supplied = clean(entry.get("entry_id"))
    if supplied:
        return supplied
    key = "|".join(
        clean(entry.get(k))
        for k in ("activity_date", "salesperson", "office_name", "summary")
    )
    return "ACT-" + hashlib.sha1(key.encode("utf-8")).hexdigest()[:12].upper()


def rows_as_dicts(ws):
    if ws.max_row < 2:
        return []
    headers = [clean(c.value) for c in ws[1]]
    result = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in row):
            result.append(dict(zip(headers, row)))
    return result


def reset_sheet(ws, headers):
    ws.delete_rows(1, ws.max_row)
    ws.append(headers)


def ensure_workbook(path):
    if path and path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    for name in SHEETS:
        if name not in wb.sheetnames:
            wb.create_sheet(name)
    if wb["사업소_목록"].max_row == 1 and wb["사업소_목록"].cell(1, 1).value is None:
        reset_sheet(wb["사업소_목록"], OFFICE_HEADERS)
    if wb["영업일지"].max_row == 1 and wb["영업일지"].cell(1, 1).value is None:
        reset_sheet(wb["영업일지"], LOG_HEADERS)
    if wb["확인필요"].max_row == 1 and wb["확인필요"].cell(1, 1).value is None:
        reset_sheet(wb["확인필요"], REVIEW_HEADERS)
    return wb


def next_office_id(office_rows):
    nums = []
    for row in office_rows:
        match = re.search(r"(\d+)$", clean(row.get("사업소ID")))
        if match:
            nums.append(int(match.group(1)))
    return f"OFF-{max(nums, default=0) + 1:03d}"


def add_review(reviews, kind, activity_id, office, detail, action):
    item = {
        "유형": kind,
        "활동ID": activity_id,
        "사업소명": office,
        "내용": detail,
        "권장조치": action,
    }
    key = tuple(clean(item[h]) for h in REVIEW_HEADERS)
    if key not in {tuple(clean(x.get(h)) for h in REVIEW_HEADERS) for x in reviews}:
        reviews.append(item)


def style_sheet(ws, widths=None):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[1].height = 26
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
    if ws.max_row >= 2:
        ws.auto_filter.ref = ws.dimensions


def add_table(ws, name):
    if ws.max_row < 2 or ws.max_column < 1:
        return
    for old in list(ws.tables.values()):
        del ws.tables[old.name]
    table = Table(displayName=name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)


def write_rows(ws, headers, rows):
    reset_sheet(ws, headers)
    for item in rows:
        ws.append([item.get(h) for h in headers])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--entries", required=True, type=Path)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    raw_entries = json.loads(args.entries.read_text(encoding="utf-8"))
    if not isinstance(raw_entries, list):
        raise SystemExit("--entries must contain a JSON array")

    wb = ensure_workbook(args.workbook)
    office_ws = wb["사업소_목록"]
    log_ws = wb["영업일지"]
    existing_offices = rows_as_dicts(office_ws)
    existing_logs = rows_as_dicts(log_ws)
    reviews = rows_as_dicts(wb["확인필요"])

    office_by_norm = {norm_name(r.get("사업소명")): r for r in existing_offices if norm_name(r.get("사업소명"))}
    existing_ids = {clean(r.get("활동ID")) for r in existing_logs}
    composite_keys = {
        (str(parse_date(r.get("활동일")) or ""), clean(r.get("담당자")), norm_name(r.get("사업소명")), clean(r.get("상담내용")))
        for r in existing_logs
    }

    inserted = 0
    duplicates = 0
    rejected = 0
    new_offices = 0

    for raw in raw_entries:
        if not isinstance(raw, dict):
            rejected += 1
            add_review(reviews, "입력형식", "", "", "객체가 아닌 항목", "원본 형식을 확인")
            continue
        activity_id = stable_id(raw)
        activity_date = parse_date(raw.get("activity_date"))
        salesperson = clean(raw.get("salesperson"))
        office_name = clean(raw.get("office_name"))
        summary = clean(raw.get("summary"))
        missing = [name for name, value in (("활동일", activity_date), ("담당자", salesperson), ("사업소명", office_name), ("상담내용", summary)) if not value]
        if missing:
            rejected += 1
            add_review(reviews, "필수값누락", activity_id, office_name, ", ".join(missing) + " 누락", "원본 일지 확인 후 재실행")
            continue
        composite = (str(activity_date), salesperson, norm_name(office_name), summary)
        if activity_id in existing_ids or composite in composite_keys:
            duplicates += 1
            continue

        office_key = norm_name(office_name)
        office = office_by_norm.get(office_key)
        if office is None:
            candidates = []
            for key, row in office_by_norm.items():
                score = SequenceMatcher(None, office_key, key).ratio()
                if score >= 0.78:
                    candidates.append((score, row))
            if candidates:
                candidates.sort(key=lambda x: x[0], reverse=True)
                candidate = candidates[0][1]
                add_review(
                    reviews, "유사사업소", activity_id, office_name,
                    f"기존 '{candidate.get('사업소명')}'와 이름이 유사함",
                    "동일 사업소인지 확인 후 병합"
                )
            office = {
                "사업소ID": next_office_id(existing_offices), "사업소명": office_name,
                "지역": clean(raw.get("region")), "담당자": salesperson,
                "진행단계": clean(raw.get("stage")), "관심품목": ", ".join(split_products(raw.get("products"))),
                "최근활동일": activity_date, "다음행동": clean(raw.get("next_action")),
                "다음행동일": parse_date(raw.get("next_action_date")), "누적활동수": 0,
                "등록일": activity_date,
            }
            existing_offices.append(office)
            office_by_norm[office_key] = office
            new_offices += 1

        next_action = clean(raw.get("next_action"))
        next_action_date = parse_date(raw.get("next_action_date"))
        if next_action and not next_action_date:
            add_review(reviews, "예정일누락", activity_id, office_name, f"'{next_action}'의 예정일 없음", "담당자에게 일정 확인")

        row = {
            "활동ID": activity_id, "활동일": activity_date, "담당자": salesperson,
            "사업소ID": office["사업소ID"], "사업소명": office_name,
            "지역": clean(raw.get("region")) or clean(office.get("지역")),
            "접촉방식": clean(raw.get("contact_type")), "상담내용": summary,
            "관심품목": ", ".join(split_products(raw.get("products"))),
            "진행단계": clean(raw.get("stage")), "다음행동": next_action,
            "예정일": next_action_date,
        }
        existing_logs.append(row)
        existing_ids.add(activity_id)
        composite_keys.add(composite)
        inserted += 1

    logs_by_office = defaultdict(list)
    for row in existing_logs:
        logs_by_office[clean(row.get("사업소ID"))].append(row)
    for office in existing_offices:
        logs = sorted(logs_by_office.get(clean(office.get("사업소ID")), []), key=lambda r: parse_date(r.get("활동일")) or date.min)
        if not logs:
            continue
        latest = logs[-1]
        products = sorted({p for row in logs for p in split_products(row.get("관심품목"))})
        office["지역"] = clean(latest.get("지역")) or office.get("지역")
        office["담당자"] = clean(latest.get("담당자")) or office.get("담당자")
        office["진행단계"] = clean(latest.get("진행단계")) or office.get("진행단계")
        office["관심품목"] = ", ".join(products)
        office["최근활동일"] = parse_date(latest.get("활동일"))
        office["누적활동수"] = len(logs)
        pending = [r for r in logs if clean(r.get("다음행동"))]
        if pending:
            target = pending[-1]
            office["다음행동"] = clean(target.get("다음행동"))
            office["다음행동일"] = parse_date(target.get("예정일"))

    write_rows(office_ws, OFFICE_HEADERS, sorted(existing_offices, key=lambda r: clean(r.get("사업소ID"))))
    write_rows(log_ws, LOG_HEADERS, sorted(existing_logs, key=lambda r: (parse_date(r.get("활동일")) or date.min, clean(r.get("활동ID")))))

    old_follow_status = {clean(r.get("활동ID")): clean(r.get("상태")) for r in rows_as_dicts(wb["후속조치"])}
    today = date.today()
    latest_day_by_office = {}
    for row in existing_logs:
        office_id = clean(row.get("사업소ID"))
        activity_day = parse_date(row.get("활동일")) or date.min
        latest_day_by_office[office_id] = max(latest_day_by_office.get(office_id, date.min), activity_day)
    follows = []
    for row in existing_logs:
        action = clean(row.get("다음행동"))
        if not action:
            continue
        due = parse_date(row.get("예정일"))
        prior = old_follow_status.get(clean(row.get("활동ID")))
        activity_day = parse_date(row.get("활동일")) or date.min
        later_activity_exists = latest_day_by_office.get(clean(row.get("사업소ID")), date.min) > activity_day
        status = prior if prior in {"완료", "보류"} else ("완료" if later_activity_exists else ("지연" if due and due < today else "예정"))
        follows.append({
            "활동ID": row.get("활동ID"), "사업소명": row.get("사업소명"), "담당자": row.get("담당자"),
            "다음행동": action, "예정일": due, "상태": status, "근거활동일": parse_date(row.get("활동일")),
        })
    follows.sort(key=lambda r: (r.get("상태") != "지연", r.get("예정일") or date.max))
    write_rows(wb["후속조치"], FOLLOW_HEADERS, follows)

    product_counts = Counter(p for row in existing_logs for p in split_products(row.get("관심품목")))
    total_mentions = sum(product_counts.values())
    product_rows = [{"품목": p, "상담건수": c, "관심도": c / total_mentions if total_mentions else 0} for p, c in product_counts.most_common()]
    write_rows(wb["품목_관심도"], ["품목", "상담건수", "관심도"], product_rows)

    valid_dates = [parse_date(r.get("활동일")) for r in existing_logs if parse_date(r.get("활동일"))]
    latest_date = max(valid_dates, default=today)
    week_start = latest_date - timedelta(days=latest_date.weekday())
    week_end = week_start + timedelta(days=6)
    week_logs = [r for r in existing_logs if (parse_date(r.get("활동일")) and week_start <= parse_date(r.get("활동일")) <= week_end)]
    stage_counts = Counter(clean(r.get("진행단계")) or "미분류" for r in week_logs)
    person_counts = Counter(clean(r.get("담당자")) or "미지정" for r in week_logs)
    ws = wb["주간_현황"]
    ws.delete_rows(1, ws.max_row)
    ws.append(["주간 영업 현황", f"{week_start.isoformat()} ~ {week_end.isoformat()}"])
    ws.append([])
    ws.append(["핵심지표", "값"])
    ws.append(["영업활동", len(week_logs)])
    ws.append(["접촉 사업소", len({clean(r.get("사업소ID")) for r in week_logs})])
    ws.append(["신규 사업소", sum(1 for r in existing_offices if week_start <= (parse_date(r.get("등록일")) or date.min) <= week_end)])
    ws.append(["후속조치", sum(1 for r in follows if r.get("상태") != "완료" and week_start <= (r.get("근거활동일") or date.min) <= week_end)])
    ws.append(["확인필요", len(reviews)])
    ws.append([])
    stage_row = ws.max_row + 1
    ws.append(["진행단계", "활동건수"])
    for key, value in stage_counts.most_common():
        ws.append([key, value])
    person_row = ws.max_row + 2
    ws.append([])
    ws.append(["담당자", "활동건수"])
    for key, value in person_counts.most_common():
        ws.append([key, value])

    write_rows(wb["확인필요"], REVIEW_HEADERS, reviews)

    style_sheet(office_ws, {"A": 12, "B": 22, "C": 15, "D": 12, "E": 14, "F": 24, "G": 13, "H": 22, "I": 13, "J": 12, "K": 13})
    style_sheet(log_ws, {"A": 22, "B": 13, "C": 12, "D": 12, "E": 22, "F": 15, "G": 12, "H": 44, "I": 24, "J": 14, "K": 22, "L": 13})
    style_sheet(wb["후속조치"], {"A": 22, "B": 22, "C": 12, "D": 28, "E": 13, "F": 10, "G": 13})
    style_sheet(wb["품목_관심도"], {"A": 22, "B": 12, "C": 12})
    style_sheet(wb["확인필요"], {"A": 16, "B": 22, "C": 22, "D": 42, "E": 30})
    for sheet, table in ((office_ws, "OfficeTable"), (log_ws, "ActivityTable"), (wb["후속조치"], "FollowupTable"), (wb["품목_관심도"], "ProductTable"), (wb["확인필요"], "ReviewTable")):
        add_table(sheet, table)

    for cell in wb["품목_관심도"]["C"][1:]:
        cell.number_format = "0.0%"
    for sheet_name in ("사업소_목록", "영업일지", "후속조치"):
        for row in wb[sheet_name].iter_rows():
            for cell in row:
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"

    follow_ws = wb["후속조치"]
    if follow_ws.max_row >= 2:
        follow_ws.conditional_formatting.add(f"A2:G{follow_ws.max_row}", FormulaRule(formula=["$F2=\"지연\""], fill=PatternFill("solid", fgColor=RED)))
        follow_ws.conditional_formatting.add(f"A2:G{follow_ws.max_row}", FormulaRule(formula=["$F2=\"완료\""], fill=PatternFill("solid", fgColor=GREEN)))

    product_ws = wb["품목_관심도"]
    product_ws._charts = []
    if product_ws.max_row >= 2:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "품목별 상담 건수"
        chart.y_axis.title = "품목"
        chart.x_axis.title = "건수"
        chart.add_data(Reference(product_ws, min_col=2, min_row=1, max_row=product_ws.max_row), titles_from_data=True)
        chart.set_categories(Reference(product_ws, min_col=1, min_row=2, max_row=product_ws.max_row))
        chart.height = 7
        chart.width = 12
        product_ws.add_chart(chart, "E2")

    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["B1"].font = Font(size=13, bold=True, color=WHITE)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=NAVY)
    for cell in (ws["A3"], ws["B3"], ws.cell(stage_row, 1), ws.cell(stage_row, 2), ws.cell(person_row, 1), ws.cell(person_row, 2)):
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
    for row in range(4, 9):
        ws.cell(row, 1).font = Font(bold=True, color=NAVY)
        ws.cell(row, 2).font = Font(bold=True, size=14, color=BLUE)

    for chart in list(ws._charts):
        ws._charts.remove(chart)
    if stage_counts:
        pie = PieChart()
        pie.title = "진행단계별 활동"
        pie.add_data(Reference(ws, min_col=2, min_row=stage_row, max_row=stage_row + len(stage_counts)), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=1, min_row=stage_row + 1, max_row=stage_row + len(stage_counts)))
        pie.height = 7
        pie.width = 10
        ws.add_chart(pie, "D3")
    if person_counts:
        bar = BarChart()
        bar.title = "담당자별 활동"
        bar.add_data(Reference(ws, min_col=2, min_row=person_row, max_row=person_row + len(person_counts)), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=person_row + 1, max_row=person_row + len(person_counts)))
        bar.height = 7
        bar.width = 10
        ws.add_chart(bar, "D18")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    pending_followups = sum(1 for row in follows if row.get("상태") != "완료")
    summary = {
        "output": str(args.output.resolve()), "inserted": inserted, "duplicates_skipped": duplicates,
        "rejected": rejected, "new_offices": new_offices, "total_activities": len(existing_logs),
        "followups": pending_followups, "followups_total": len(follows), "review_items": len(reviews),
        "week": f"{week_start.isoformat()}~{week_end.isoformat()}",
    }
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
