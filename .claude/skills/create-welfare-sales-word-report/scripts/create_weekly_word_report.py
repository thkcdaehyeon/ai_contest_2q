#!/usr/bin/env python3
"""Generate a Korean weekly sales report DOCX from the linked workbook."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import load_workbook


NAVY = "17324D"
BLUE = "2563EB"
CYAN = "0EA5A8"
PALE = "EAF2F8"
LIGHT = "F8FAFC"
RED = "B91C1C"
GRAY = "64748B"
FONT = "Arial Unicode MS"


def clean(value):
    return str(value or "").strip()


def parse_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None


def rows_as_dicts(ws):
    headers = [clean(c.value) for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in row)]


def split_products(value):
    return [x.strip() for x in clean(value).replace(";", ",").split(",") if x.strip()]


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text, bold=False, color="17324D", size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(clean(text))
    r.bold = bold
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(17 if level == 1 else 12)
    r.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)
    return p


def add_note(doc, text, color=GRAY):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_table(doc, headers, data, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, header in enumerate(headers):
        set_cell_fill(table.rows[0].cells[i], NAVY)
        set_cell_text(table.rows[0].cells[i], header, True, "FFFFFF", 8.5, WD_ALIGN_PARAGRAPH.CENTER)
    for row_idx, row in enumerate(data):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, NAVY, 8.2)
            if row_idx % 2:
                set_cell_fill(cells[i], LIGHT)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--week-start")
    parser.add_argument("--week-end")
    args = parser.parse_args()

    wb = load_workbook(args.workbook, data_only=True)
    if "영업일지" not in wb.sheetnames:
        raise SystemExit("Workbook must contain 영업일지")
    logs = rows_as_dicts(wb["영업일지"])
    dated = [parse_date(r.get("활동일")) for r in logs if parse_date(r.get("활동일"))]
    if not dated:
        raise SystemExit("영업일지 contains no valid activity dates")
    if args.week_start:
        week_start = parse_date(args.week_start)
        if not week_start:
            raise SystemExit("Invalid --week-start")
    else:
        latest = max(dated)
        week_start = latest - timedelta(days=latest.weekday())
    if args.week_end:
        week_end = parse_date(args.week_end)
        if not week_end:
            raise SystemExit("Invalid --week-end")
    else:
        week_end = week_start + timedelta(days=6)
    if week_end < week_start:
        raise SystemExit("--week-end must be on or after --week-start")

    week_logs = [r for r in logs if parse_date(r.get("활동일")) and week_start <= parse_date(r.get("활동일")) <= week_end]
    week_logs.sort(key=lambda r: (parse_date(r.get("활동일")), clean(r.get("담당자")), clean(r.get("사업소명"))))
    ids = {clean(r.get("활동ID")) for r in week_logs}
    follows = rows_as_dicts(wb["후속조치"]) if "후속조치" in wb.sheetnames else []
    follows = [r for r in follows if clean(r.get("활동ID")) in ids and clean(r.get("상태")) != "완료"]
    reviews = rows_as_dicts(wb["확인필요"]) if "확인필요" in wb.sheetnames else []
    reviews = [r for r in reviews if not clean(r.get("활동ID")) or clean(r.get("활동ID")) in ids]

    product_counts = Counter(p for r in week_logs for p in split_products(r.get("관심품목")))
    stage_counts = Counter(clean(r.get("진행단계")) or "미분류" for r in week_logs)
    person_counts = Counter(clean(r.get("담당자")) or "미지정" for r in week_logs)
    unique_offices = len({clean(r.get("사업소ID")) or clean(r.get("사업소명")) for r in week_logs})

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Normal"].font.size = Pt(9.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("주간 영업 활동 보고서")
    r.bold = True
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"{week_start.isoformat()} － {week_end.isoformat()}")
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor.from_string(GRAY)

    kpis = [
        ("영업활동", f"{len(week_logs)}건"),
        ("접촉 사업소", f"{unique_offices}곳"),
        ("후속조치", f"{len(follows)}건"),
        ("확인필요", f"{len(reviews)}건"),
    ]
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(kpis):
        set_cell_fill(table.cell(0, i), PALE)
        set_cell_fill(table.cell(1, i), PALE)
        set_cell_text(table.cell(0, i), label, True, GRAY, 9, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(1, i), value, True, BLUE, 17, WD_ALIGN_PARAGRAPH.CENTER)

    add_heading(doc, "이번 주 요약")
    if week_logs:
        top_product = product_counts.most_common(1)[0] if product_counts else None
        top_stage = stage_counts.most_common(1)[0] if stage_counts else None
        sentences = [f"총 {len(week_logs)}건의 영업 활동으로 {unique_offices}개 사업소와 접촉했습니다."]
        if top_product:
            sentences.append(f"가장 많이 언급된 관심 품목은 {top_product[0]}({top_product[1]}건)입니다.")
        if top_stage:
            sentences.append(f"가장 많은 진행 단계는 {top_stage[0]}({top_stage[1]}건)입니다.")
        sentences.append(f"미완료 후속 조치는 {len(follows)}건이며, 사람이 확인할 데이터는 {len(reviews)}건입니다.")
        add_note(doc, " ".join(sentences), NAVY)
    else:
        add_note(doc, "선택한 기간에 영업 활동이 없습니다.", RED)

    add_heading(doc, "품목 관심도와 활동 현황")
    product_total = sum(product_counts.values())
    product_data = [[p, f"{c}건", f"{(c / product_total * 100):.1f}%" if product_total else "0.0%"] for p, c in product_counts.most_common()]
    add_table(doc, ["품목", "상담 건수", "관심도"], product_data or [["기록 없음", "0건", "0.0%"]], [8, 3.5, 3.5])
    stage_data = [[name, f"{count}건"] for name, count in stage_counts.most_common()]
    person_data = [[name, f"{count}건"] for name, count in person_counts.most_common()]
    pair_count = max(len(stage_data), len(person_data), 1)
    paired = []
    for i in range(pair_count):
        left = stage_data[i] if i < len(stage_data) else ["", ""]
        right = person_data[i] if i < len(person_data) else ["", ""]
        paired.append(left + right)
    add_table(doc, ["진행 단계", "건수", "담당자", "활동 건수"], paired, [5.5, 2.5, 5.5, 2.5])

    add_heading(doc, "주요 영업 활동")
    activity_data = []
    for r in week_logs:
        d = parse_date(r.get("활동일"))
        activity_data.append([
            d.strftime("%m/%d") if d else "", clean(r.get("담당자")), clean(r.get("사업소명")),
            clean(r.get("진행단계")) or "미분류", clean(r.get("상담내용")), clean(r.get("다음행동")) or "－",
        ])
    add_table(doc, ["일자", "담당", "사업소", "단계", "상담 내용", "다음 행동"], activity_data or [["", "", "기록 없음", "", "", ""]], [1.6, 2, 3.2, 2.2, 6.3, 3.2])

    add_heading(doc, "후속 조치")
    follow_data = []
    for r in sorted(follows, key=lambda x: parse_date(x.get("예정일")) or date.max):
        due = parse_date(r.get("예정일"))
        follow_data.append([
            clean(r.get("상태")), due.isoformat() if due else "일정 확인 필요", clean(r.get("담당자")),
            clean(r.get("사업소명")), clean(r.get("다음행동")),
        ])
    add_table(doc, ["상태", "예정일", "담당", "사업소", "할 일"], follow_data or [["－", "－", "－", "미완료 후속 조치 없음", "－"]], [2, 3, 2.2, 4.2, 7])

    add_heading(doc, "확인 필요")
    review_data = [[clean(r.get("유형")), clean(r.get("사업소명")), clean(r.get("내용")), clean(r.get("권장조치"))] for r in reviews]
    add_table(doc, ["유형", "사업소", "확인 내용", "권장 조치"], review_data or [["없음", "－", "확인할 예외가 없습니다.", "－"]], [2.7, 4, 7, 5])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("영업관리 Excel 기반 자동 생성 · 수치는 원본 대장과 함께 검증")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(GRAY)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(json.dumps({
        "output": str(args.output.resolve()), "week": f"{week_start.isoformat()}~{week_end.isoformat()}",
        "activities": len(week_logs), "offices": unique_offices, "followups": len(follows),
        "review_items": len(reviews), "products": len(product_counts),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
